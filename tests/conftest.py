import pytest
import os
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter

@pytest.fixture(scope="session")
def test_data_dir(tmpdir_factory):
    """Create a temporary directory for test data."""
    return tmpdir_factory.mktemp("data")

@pytest.fixture(scope="session")
def test_pdf_path(test_data_dir):
    """Creates a dummy PDF file for testing and returns its path."""
    pdf_path = os.path.join(test_data_dir, "test_document.pdf")
    c = canvas.Canvas(str(pdf_path), pagesize=letter)
    c.drawString(72, 800, "The sky is blue. This is a test document for PDF extraction.")
    c.save()
    return str(pdf_path)

@pytest.fixture(scope="session")
def test_docx_path(test_data_dir):
    """Creates a dummy DOCX file for testing and returns its path."""
    docx_path = os.path.join(test_data_dir, "test_document.docx")
    document = Document()
    document.add_paragraph("The grass is green. This is a test document for DOCX extraction.")
    document.save(docx_path)
    return str(docx_path)

@pytest.fixture(scope="session")
def empty_docx_path(test_data_dir):
    """Creates an empty DOCX file."""
    docx_path = os.path.join(test_data_dir, "empty.docx")
    document = Document()
    document.save(docx_path)
    return str(docx_path)